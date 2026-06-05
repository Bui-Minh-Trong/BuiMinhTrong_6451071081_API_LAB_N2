from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports"
OUT_FILE = OUT_DIR / "BaoCao_Bai1_TichHopFacebookAPI_Backend.docx"


COLORS = {
    "blue": RGBColor(46, 116, 181),
    "dark_blue": RGBColor(31, 77, 120),
    "ink": RGBColor(30, 41, 59),
    "muted": RGBColor(100, 116, 139),
    "light_fill": "F2F4F7",
    "callout_fill": "F4F6F9",
    "border": "D0D7DE",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(row.cells[idx])


def style_table(table, widths):
    table.style = "Table Grid"
    set_table_width(table, widths)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, COLORS["light_fill"])
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = COLORS["ink"]


def add_run(paragraph, text, bold=False, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return run


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.style = "Code"
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9.5)


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(6.35)
    cell = table.cell(0, 0)
    cell.width = Inches(6.35)
    set_cell_shading(cell, COLORS["callout_fill"])
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_run(p, title + ": ", bold=True, color=COLORS["dark_blue"])
    add_run(p, text)
    doc.add_paragraph()


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = COLORS["ink"]
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Heading 1", 16, COLORS["blue"], 16, 8),
        ("Heading 2", 13, COLORS["blue"], 12, 6),
        ("Heading 3", 12, COLORS["dark_blue"], 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    if "Code" not in styles:
        code = styles.add_style("Code", 1)
    else:
        code = styles["Code"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(9.5)
    code.paragraph_format.space_after = Pt(0)


def write_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = "Báo cáo thực hành Lập trình API - Bài 1"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.size = Pt(9)
    hp.runs[0].font.color.rgb = COLORS["muted"]

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "Facebook Page API & Backend Proxy"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.runs[0].font.size = Pt(9)
    fp.runs[0].font.color.rgb = COLORS["muted"]


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("BÁO CÁO THỰC HÀNH LẬP TRÌNH API")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = COLORS["dark_blue"]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Bài 1: Tích hợp Facebook API và xây dựng Backend")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = COLORS["blue"]

    meta = [
        ("Đề tài", "Hệ thống quản lý Facebook Page phân tán"),
        ("Sinh vien", "Bui Minh Trong"),
        ("MSSV", "6451071081"),
        ("Workspace", "BuiMinhTrong_6451071081_LAB_N2"),
        ("Ngày lập báo cáo", "05/06/2026"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    style_table(table, [Inches(1.7), Inches(4.4)])
    for i, (key, value) in enumerate(meta):
        table.cell(i, 0).text = key
        table.cell(i, 1).text = value
        if i > 0:
            set_cell_shading(table.cell(i, 0), "FFFFFF")
            set_cell_shading(table.cell(i, 1), "FFFFFF")
        for p in table.cell(i, 0).paragraphs:
            for r in p.runs:
                r.bold = True
    doc.add_paragraph()
    add_callout(
        doc,
        "Tóm tắt",
        "Bài 1 xây dựng backend-api đóng vai trò proxy giữa dashboard/Swagger và Facebook Graph API. Client đăng nhập lấy JWT, sau đó gọi backend để đọc bài viết, đăng bài và đọc bình luận trên Page.",
    )
    doc.add_page_break()


def build_doc():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    setup_styles(doc)
    write_header_footer(doc)
    add_cover(doc)

    doc.add_heading("1. Mục tiêu bài 1", level=1)
    doc.add_paragraph(
        "Theo đề bài, Bài 1 yêu cầu tích hợp Facebook Graph API và xây dựng một backend trung gian. Swagger hoặc dashboard không gọi trực tiếp Facebook, mà gọi backend của hệ thống; backend đọc cấu hình PAGE_ID và PAGE_ACCESS_TOKEN trong file .env rồi gọi Graph API tương ứng."
    )
    add_bullets(
        doc,
        [
            "Tạo Facebook Page và Facebook App để lấy Page Access Token.",
            "Xây dựng backend-api chạy tại port 3000.",
            "Cài đặt các API chính: GET /posts, POST /post, GET /comments.",
            "Bổ sung xác thực admin bằng JWT cho các API quản trị.",
            "Chuẩn hóa response, mã lỗi và log khi gọi Facebook Graph API.",
        ],
    )

    doc.add_heading("2. Cấu hình và thành phần liên quan", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "Thành phần"
    table.rows[0].cells[1].text = "Port/File"
    table.rows[0].cells[2].text = "Vai trò"
    table.rows[0].cells[3].text = "Ghi chú demo"
    rows = [
        ("backend-api", "3000", "Expose REST API và Swagger UI", "Service chính của Bài 1"),
        ("Swagger UI", "/api-docs", "Giao diện demo và test API", "Mở tại http://localhost:3000/api-docs"),
        ("Facebook Graph API", "v19.0", "Đọc post, tạo post, đọc comment", "Backend mới được gọi trực tiếp"),
        ("backend-api/.env", "PAGE_ID, PAGE_ACCESS_TOKEN, JWT_SECRET", "Lưu cấu hình bí mật", "Không public token trong báo cáo"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, [Inches(1.35), Inches(1.55), Inches(2.0), Inches(2.05)])

    doc.add_heading("3. Các thao tác thực hiện", level=1)
    add_steps(
        doc,
        [
            "Tạo Facebook Page dùng để demo quản lý bài viết và bình luận.",
            "Tạo Meta Developer App và lấy Page Access Token có quyền đọc Page, đăng bài và đọc engagement.",
            "Cấu hình backend-api/.env với PORT=3000, PAGE_ID, PAGE_ACCESS_TOKEN, JWT_SECRET, ADMIN_USER và ADMIN_PASS.",
            "Cài đặt dependency cho backend-api bằng npm install trong thư mục backend-api.",
            "Chạy backend-api độc lập cho Bài 1 bằng npm run demo:bai1 hoặc npm run start:backend tại thư mục gốc project.",
            "Mở Swagger UI tại http://localhost:3000/api-docs.",
            "Đăng nhập bằng POST /auth/login, copy access_token và nhập vào nút Authorize của Swagger.",
            "Demo lần lượt GET /posts, POST /post và GET /comments?post_id=<post_id>.",
        ],
    )

    doc.add_heading("4. Cấu trúc API đã cài đặt", level=1)
    table = doc.add_table(rows=1, cols=5)
    headers = ["API", "Method", "Input chính", "Xử lý backend", "Kết quả"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    api_rows = [
        ("Đăng nhập admin", "POST /auth/login", "username, password", "Kiểm tra ADMIN_USER/ADMIN_PASS và ký JWT", "Trả access_token Bearer"),
        ("Lấy bài viết Page", "GET /posts", "Bearer token, page_id tùy chọn", "Gọi Graph API /{pageId}/posts", "Danh sách bài viết"),
        ("Đăng bài viết", "POST /post", "Bearer token, message", "Validate message, gọi /{pageId}/feed", "ID bài viết mới"),
        ("Lấy bình luận", "GET /comments", "Bearer token, post_id", "Gọi /{postId}/comments", "Danh sách comment"),
        ("Reply comment", "POST /comments/:commentId/reply", "message", "Gọi /{commentId}/comments", "Phản hồi admin"),
        ("Ẩn comment", "POST /comments/:commentId/hide", "commentId", "Gọi /{commentId} với is_hidden=true", "Comment được ẩn"),
    ]
    for row in api_rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, [Inches(1.35), Inches(1.35), Inches(1.45), Inches(2.0), Inches(1.0)])

    doc.add_heading("5. Định dạng response và xử lý lỗi", level=1)
    doc.add_paragraph(
        "Backend trả response thống nhất theo format success/data/error/timestamp. Khi thiếu input, token sai, token hết hạn hoặc Facebook Graph API lỗi, service trả mã lỗi rõ ràng thay vì để client nhận lỗi thấp tầng."
    )
    add_code_block(
        doc,
        [
            "Thanh cong:",
            '{ "success": true, "data": { ... }, "error": null, "timestamp": "2026-06-05T00:00:00.000Z" }',
            "",
            "That bai:",
            '{ "success": false, "data": null, "error": { "code": "FB_UNAUTHORIZED", "message": "Facebook API rejected the page token or permissions." }, "timestamp": "..." }',
        ],
    )
    add_callout(
        doc,
        "Ghi chú vận hành",
        "Nếu Facebook trả lỗi (#200) Permissions error, Kafka và backend không sai; cần kiểm tra lại Page Access Token và các quyền pages_read_engagement/pages_manage_engagement.",
    )

    doc.add_heading("6. Kịch bản demo Bài 1", level=1)
    table = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["Bước", "Thao tác demo", "Dữ liệu minh chứng", "Kết quả mong đợi"]):
        table.rows[0].cells[i].text = h
    demo_rows = [
        ("1", "Mở backend-api", "http://localhost:3000/health", "status=ok, service=backend-api"),
        ("2", "Mở Swagger", "http://localhost:3000/api-docs", "Hiện danh sách endpoint"),
        ("3", "Login", "POST /auth/login với admin/admin123", "Nhận token_type=Bearer"),
        ("4", "Lấy bài viết", "GET /posts", "Trả về post trên Page thật"),
        ("5", "Đăng bài", "POST /post với message demo", "Facebook Page xuất hiện bài viết mới"),
        ("6", "Lấy comment", "GET /comments?post_id=<post_id>", "Trả về danh sách comment của bài viết"),
    ]
    for row in demo_rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, [Inches(0.55), Inches(1.65), Inches(2.35), Inches(2.35)])

    doc.add_heading("7. Kết quả demo đã ghi nhận", level=1)
    doc.add_paragraph(
        "Tại thời điểm kiểm tra 05/06/2026, backend-api đang chạy và các endpoint đọc dữ liệu Facebook của Bài 1 hoạt động. Kết quả tóm tắt:"
    )
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Hạng mục", "Kết quả", "Minh chứng"]):
        table.rows[0].cells[i].text = h
    result_rows = [
        ("Health backend", "Thành công", "GET /health trả service=backend-api, database_enabled=true, kafka_enabled=true"),
        ("Login JWT", "Thành công", "POST /auth/login trả success=true và token_type=Bearer"),
        ("GET /posts", "Thành công", "Lấy được bài viết: 'Mua gì bán nấy' và 'Demo Bai 1 - dang bai qua backend proxy'"),
        ("GET /comments", "Thành công", "Lấy được comment: 'giá sao vậy shop ?' và 'Toi dang test get / comments'"),
    ]
    for row in result_rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, [Inches(1.35), Inches(1.15), Inches(4.4)])

    doc.add_heading("8. Lệnh chạy và lệnh kiểm thử mẫu", level=1)
    add_code_block(
        doc,
        [
            'cd "D:\\UTC2\\Nam3\\HK2_2025-2026\\API\\BuiMinhTrong_6451071081_LAB_N2"',
            "npm run start:backend",
            "",
            'curl -X POST http://localhost:3000/auth/login -H "Content-Type: application/json" -d "{\\"username\\":\\"admin\\",\\"password\\":\\"admin123\\"}"',
            'curl http://localhost:3000/posts -H "Authorization: Bearer <access_token>"',
            'curl -X POST http://localhost:3000/post -H "Authorization: Bearer <access_token>" -H "Content-Type: application/json" -d "{\\"message\\":\\"Demo Bai 1 - dang bai qua backend proxy\\"}"',
            'curl "http://localhost:3000/comments?post_id=<post_id>" -H "Authorization: Bearer <access_token>"',
        ],
    )

    doc.add_heading("9. Đánh giá đáp ứng yêu cầu", level=1)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Yêu cầu đề bài", "Trạng thái", "Giải thích"]):
        table.rows[0].cells[i].text = h
    checklist = [
        ("Backend proxy Facebook API", "Đạt", "Client chỉ gọi backend-api; Graph API được gọi trong services/facebookApi.js."),
        ("GET /posts", "Đạt", "Endpoint có auth, lấy page_id từ query hoặc .env."),
        ("POST /post", "Đạt", "Có validate message và trả response 201 khi thành công."),
        ("GET /comments", "Đạt", "Nhận post_id và gọi Graph API lấy comment."),
        ("Swagger UI", "Đạt", "Chạy tại /api-docs để demo trực tiếp."),
        ("JWT admin", "Đạt", "POST /auth/login cấp token, route quản trị yêu cầu Bearer token."),
        ("Response/error chuẩn hóa", "Đạt", "Các route trả success/data/error/timestamp; lỗi Facebook được map thành code rõ ràng."),
        ("Log request Facebook", "Đạt", "facebookApi.js ghi method, endpoint, status và latency."),
    ]
    for row in checklist:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, [Inches(2.2), Inches(0.9), Inches(3.85)])

    doc.add_heading("10. Kết luận", level=1)
    doc.add_paragraph(
        "Bài 1 đã hoàn thành mục tiêu xây dựng backend-api làm lớp trung gian với Facebook Graph API. Hệ thống có Swagger để demo, có JWT cho dashboard quản trị, có các API chính để đọc bài viết, đăng bài và đọc bình luận, đồng thời chuẩn hóa response và xử lý lỗi rõ ràng. Đây là nền tảng để sang Bài 2 kết nối webhook-service, Kafka và core-service theo kiến trúc event-driven."
    )

    doc.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    out = build_doc()
    print(out)
